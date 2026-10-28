"""保存Excelを読み取り専用で検査し、元会計・正式表・数式・対象範囲を照合する。"""

from pathlib import Path
from datetime import datetime, timezone, timedelta
import hashlib
import json
import zipfile
import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
BOOK = ROOT / "アルゴリズム調査/00159_支持スコアの確度検証と夜間配分/支持スコアと昼夜の全セッション損益.xlsx"


def digest(path):
    """ファイルのSHA256を求める。

    引数: path(Path)。返り値: str。
    """
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    """保存Excelの値・式と母資料を照合する。Excelは書き換えない。

    引数: なし。返り値: None。
    """
    report_path = ROOT / "work/research159/report159_all.json"
    report = json.loads(report_path.read_text(encoding="utf-8"))
    qa = json.loads((HERE / "xlsx_checks/Excel照合159.json").read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(BOOK, data_only=False)
    cached = openpyxl.load_workbook(BOOK, data_only=True)
    checked = 0

    def equal(actual, expected, label):
        """数値又は文字列を比較する。

        引数: actual/expected(object)、label(str)。返り値: None。不一致なら停止。
        """
        nonlocal checked
        checked += 1
        if isinstance(expected,(int,float)):
            assert isinstance(actual,(int,float)) and abs(actual-expected)<1e-5,(label,actual,expected)
        else:
            assert actual == expected,(label,actual,expected)

    equal(wb.sheetnames,["比較","支持の確度"]+["日経_"+m["id"] for m in report["methods"]],"シート一覧")
    for name,address,expected in qa["check_cells"]:
        equal(cached[name][address].value,expected,f"{name}/{address}")
    session_rows = 0
    for mi,m in enumerate(report["methods"]):
        name="日経_"+m["id"]
        last=len(m["rows"])+5
        total=last+1
        end="L" if m["kind"]=="Night" else "J"
        s=wb[name]
        equal(s.freeze_panes,"A6",name+"/見出し固定")
        equal(len(s.tables),1,name+"/正式表")
        table=s.tables[list(s.tables)[0]]
        equal(table.ref,f"A5:{end}{last}",name+"/表範囲")
        assert table.autoFilter is not None
        cumulative=0
        for i,row in enumerate(m["rows"]):
            session_rows += 1
            n=i+6
            cumulative+=row["net"]
            equal(cached[name][f"A{n}"].value,row["session"],name+"/対象場")
            equal(cached[name][f"B{n}"].value,row.get("sample_group","旧完了場"),name+"/資料区分")
            equal(cached[name][f"C{n}"].value,row["direction"],name+"/方向")
            equal(s[f"F{n}"].value,f"=D{n}-E{n}",name+"/収支式")
            equal(s[f"G{n}"].value,f"=F{n}" if i==0 else f"=G{n-1}+F{n}",name+"/累計式")
            for col,value in [("D",row["profit"]),("E",row["loss"]),("F",row["net"]),("G",cumulative),
                              ("H",row["intraday_risk"]["min_equity"]),("I",row["intraday_risk"]["max_drawdown"]),("J",row["entries"])]:
                equal(cached[name][f"{col}{n}"].value,value,f"{name}/{col}{n}/元会計")
            legacy=m["kind"]=="Day" and m["id"] in ["S01","J02"] and len(m["rows"])==26 and row["sample_group"]=="旧完了場"
            assert not row.get("partial",False) and (legacy or row.get("remaining_qty")==0)
        equal(wb["比較"][f"F{mi+6}"].value,f"='{name}'!F{total}",name+"/比較参照")
        if m["id"] in ["S01","J02"]:
            equal(len(m["rows"]),26,name+"/旧Day限定")
        if "完了確認" in m["id"]:
            equal(len(m["rows"]),1,name+"/追加完了は別集計")
    formulas=0
    for s in wb:
        equal(len(s.tables),2 if s.title=="比較" else 1,s.title+"/正式表数")
        for row in s:
            for cell in row:
                if cell.data_type=="f":
                    formulas+=1
                    v=cached[s.title][cell.coordinate]
                    assert v.data_type!="e",(s.title,cell.coordinate,v.value)
                    assert isinstance(v.value,(int,float)) or (s.title=="支持の確度" and v.value in [None,""]),(s.title,cell.coordinate,v.value)
                assert cell.data_type!="e",(s.title,cell.coordinate,cell.value)
    with zipfile.ZipFile(BOOK) as z:
        tables=[n for n in z.namelist() if n.startswith("xl/tables/table") and n.endswith(".xml")]
        equal(len(tables),len(wb.sheetnames)+1,"ZIP正式表数")
    equal(digest(BOOK),qa["book_sha256"],"保存Excel指紋")
    equal(digest(report_path),qa["report_sha256"],"集計根拠指紋")
    equal(digest(BOOK),digest(ROOT/"outputs/00159_support"/BOOK.name),"公開複製")
    for source in qa["score_sources"]:
        equal(digest(Path(source["path"])),source["sha256"],"支持根拠指紋")
    result={"passed":True,"checked_at":datetime.now(timezone(timedelta(hours=9))).isoformat(),
            "checks":checked,"formula_count":formulas,"sheets":len(wb.sheetnames),"tables":len(tables),
            "session_rows":session_rows,"book_sha256":digest(BOOK),"report_sha256":digest(report_path),
            "book_modified":False,"financial_replays":0,"native_excel_recalculation":False}
    (HERE/"xlsx_checks/保存Excelの独立検算159.json").write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(result,ensure_ascii=False))


if __name__=="__main__":
    main()

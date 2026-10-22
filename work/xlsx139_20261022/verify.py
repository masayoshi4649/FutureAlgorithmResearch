"""保存済み金融結果とxlsxの全行を読み取り専用で照合する。"""

from pathlib import Path
import json
import hashlib
import openpyxl


def main() -> None:
    """全明細と集計を照合する。

    引数: なし。
    返り値: None。照合記録をJSONへ保存する。
    """
    root = Path(__file__).resolve().parents[2]
    public = root / "アルゴリズム調査/00139_複数追撃と早期参加の全体検証"
    target = public / "段階投入と早期参加の全セッション損益.xlsx"
    source = {}
    for name in ("results", "extra_results", "early_results", "recycle_results", "protect_results", "guard_results"):
        rows = json.loads((root / f"work/tranche139_20261022/{name}/結果.json").read_text(encoding="utf-8"))["rows"]
        for row in rows:
            key = (row["symbol"], row["method"], row["session"])
            assert key not in source
            source[key] = row
    workbook = openpyxl.load_workbook(target, data_only=True)
    actual_keys = set()
    for symbol, name in (("NK225", "日経セッション"), ("TOPIX", "TOPIXセッション")):
        totals = {}
        for row in workbook[name].iter_rows(min_row=7, values_only=True):
            method, date, kind, profit, loss, net, cumulative, fee, entries, quantity, held, cohort = row
            session = date.strftime("%Y-%m-%d") + ("_Day" if kind == "デイ" else "_Night")
            key = (symbol, method, session)
            assert key not in actual_keys
            actual_keys.add(key)
            expected = source[key]
            for column, value in (("profit", profit), ("loss", loss), ("net", net), ("fees", fee), ("entries", entries), ("quantity", quantity), ("max_per_side", held)):
                assert abs(expected[column] - value) < 0.001, (key, column)
            totals[method] = totals.get(method, 0) + net
            assert abs(totals[method] - cumulative) < 0.001
            assert cohort == ("既知45場" if session <= "2026-10-21_Day" else "追加2場")
    assert actual_keys == set(source)

    # ----------------------------------------
    # 書き出し後の数式・表・表示形式も確認する。
    formulas = openpyxl.load_workbook(target, data_only=False)
    for name in ("日経セッション", "TOPIXセッション"):
        sheet = formulas[name]
        assert sheet.freeze_panes == "D7"
        assert len(sheet.tables) == 1
        assert sheet["B7"].number_format == "yyyy/mm/dd"
        for number in range(7, 6869):
            assert sheet[f"F{number}"].value == f"=D{number}-E{number}"
    summary = json.loads((public / "比較集計.json").read_text(encoding="utf-8"))["summary"]
    for row in workbook["比較"].iter_rows(min_row=25, max_row=316, values_only=True):
        symbol = "NK225" if row[0] == "日経mini" else "TOPIX"
        stat = summary[symbol][row[1]]["全47場"]
        for key, value in zip(("profit", "loss", "net", "reds", "mean_loss", "max_loss", "top3_removed", "fees", "buy", "sell", "max_per_side", "no_trade"), row[3:]):
            assert abs(stat[key] - value) < 0.001, (symbol, row[1], key)
    result = {"対象": str(target), "会計数": len(actual_keys), "明細・累計・集計の全数一致": True, "日付型・数式・フィルター・固定枠": True, "ファイルSHA256": hashlib.sha256(target.read_bytes()).hexdigest(), "注記": "保存xlsxを読み戻して照合。Excelアプリでの操作テストと金融再実行は行っていない。"}
    (Path(__file__).parent / "保存xlsx照合.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

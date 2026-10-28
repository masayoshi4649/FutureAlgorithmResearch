"""完成したExcelを読み取り専用で検査し、数値と表定義の証跡を残す。"""
from pathlib import Path
from collections import Counter
import hashlib
import json
import math
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from PIL import Image, ImageDraw, ImageFont

WORK = Path(__file__).resolve().parent
ROOT = WORK.parents[1]
QA = WORK / "xlsx_checks"
NAME = "D07の利益保護と段階投入の全セッション損益.xlsx"
BOOK = ROOT / "アルゴリズム調査/00154_D07の利益返しと多枚数新規の改善" / NAME
COPY = ROOT / "outputs/00154_d07_profit" / NAME
issues = []
counts = Counter()


def check(actual, expected, label):
    """保存値と独立した期待値を比較する。

    引数: actual/expected(object)、label(str)。返り値: なし。不一致を記録する。
    """
    counts["検査数"] += 1
    okay = actual == expected
    if isinstance(expected, (float, int)) and not isinstance(expected, bool):
        okay = isinstance(actual, (float, int)) and math.isfinite(actual) and abs(actual - expected) <= 1e-5
    if not okay:
        issues.append({"対象": label, "保存値": actual, "期待": expected})


def sha(path):
    """対象ファイルの内容指紋を求める。

    引数: path(Path)。返り値: str型のSHA256。
    """
    return hashlib.sha256(path.read_bytes()).hexdigest()


def cell(sheet, address, expected):
    """保存済み数式計算値または文字を照合する。

    引数: sheet/address(str)、expected(object)。返り値: なし。
    """
    check(values[sheet][address].value, expected, f"{sheet}!{address}")


def formula(sheet, address, expected):
    """保存数式を完全一致で照合する。

    引数: sheet/address/expected(str)。返り値: なし。
    """
    check(formulas[sheet][address].value, expected, f"{sheet}!{address}の式")
    counts["数式"] += 1


before = sha(BOOK)
report = json.loads((WORK / "report154_combined.json").read_text(encoding="utf-8"))
values = openpyxl.load_workbook(BOOK, data_only=True)
formulas = openpyxl.load_workbook(BOOK, data_only=False)
methods = report["methods"]
names = ["比較"] + [f"日経_{m['id']}" for m in methods] + ["10月27日デイ途中"]
check(values.sheetnames, names, "全25シートの順序")
check(len(methods), 23, "手法数")
base = {r["session"]: r for r in next(m for m in methods if m["id"] == "D07")["rows"]}
expected_tables = {}

# ----------------------------------------
# 全51場の会計と、合計・累計・基準差を照合する。
for method in methods:
    sheet = f"日経_{method['id']}"
    rows = sorted(method["rows"], key=lambda r: r["session"])
    check(len(rows), 51, f"{sheet}の場数")
    check(len({r["session"] for r in rows}), 51, f"{sheet}の場名一意性")
    accumulated = 0
    for n, row in enumerate(rows, 6):
        accumulated += row["net"]
        expected = [row["session"], "デイ" if row["kind"] == "Day" else "ナイト", row["direction"],
                    row["profit"], row["loss"], row["net"], accumulated, row["net"]-base[row["session"]]["net"],
                    row["initial"]["net"], row["addition"]["net"], row["reentry"]["net"], row["entries"], row["entry_quantity"]]
        for col, x in enumerate(expected, 1):
            cell(sheet, f"{openpyxl.utils.get_column_letter(col)}{n}", x)
        check(sum(expected[8:11]), row["net"], f"{sheet}/{row['session']}玉分類の合計")
        formula(sheet, f"F{n}", f"=D{n}-E{n}")
        formula(sheet, f"G{n}", f"=F{n}" if n == 6 else f"=G{n-1}+F{n}")
        formula(sheet, f"H{n}", f"=F{n}-F{n}" if method["id"] == "D07" else f"=F{n}-'日経_D07'!F{n}")
    total = method["summary"]
    for col, field in {"D":"profit", "E":"loss", "F":"net", "G":"net", "I":"initial_net", "J":"addition_net", "K":"reentry_net", "L":"entries"}.items():
        cell(sheet, f"{col}57", total[field])
    cell(sheet, "H57", total["net"] - sum(r["net"] for r in base.values()))
    cell(sheet, "M57", sum(r["entry_quantity"] for r in rows))
    for col in ["D", "E", "H", "I", "J", "K", "L", "M"]:
        formula(sheet, f"{col}57", f"=SUM({col}6:{col}56)")
    formula(sheet, "F57", "=D57-E57")
    formula(sheet, "G57", "=G56")
    expected_tables[sheet] = (f"D154_{method['id']}", "A5:M56")
    counts["完了場の明細行"] += len(rows)

# ----------------------------------------
# 比較表と途中結果を、完結親・未完親の区分を保って照合する。
for n, m in enumerate(methods, 6):
    s = m["summary"]
    for col, x in {"A":m["id"], "B":m["description"], "C":s["profit"], "D":s["loss"], "E":s["net"],
                   "F":s["pf"] if s["pf"] is not None else "", "G":s["losses"], "H":s["avgloss"], "I":s["maxloss"],
                   "J":s["ex3"], "K":m["day"]["net"], "L":m["night"]["net"], "M":m["up"]["net"], "N":m["down"]["net"]}.items():
        cell("比較", f"{col}{n}", x)
    p = f"'日経_{m['id']}'!"
    net, kind, direction = f"{p}$F$6:$F$56", f"{p}$B$6:$B$56", f"{p}$C$6:$C$56"
    expected = [f"={p}D57", f"={p}E57", f"={p}F57", f'=IF(D{n}=0,"",C{n}/D{n})', f'=COUNTIFS({net},"<0")',
                f'=IF(G{n}=0,0,-SUMIFS({net},{net},"<0")/G{n})', f"=MAX(0,-MIN({net}))",
                f"=E{n}-MAX(0,LARGE({net},1))-MAX(0,LARGE({net},2))-MAX(0,LARGE({net},3))",
                f'=SUMIFS({net},{kind},"デイ")', f'=SUMIFS({net},{kind},"ナイト")',
                f'=SUMIFS({net},{direction},"上昇")', f'=SUMIFS({net},{direction},"下落")']
    for col, x in enumerate(expected, 3):
        formula("比較", f"{openpyxl.utils.get_column_letter(col)}{n}", x)
expected_tables["比較"] = ("D154_comparison", "A5:N28")
for n, r in enumerate(report["current_rows"], 6):
    stop = r["at_actual_stop"]
    expected = [r["id"],r["profit"],r["loss"],r["open_parent_realized"],r["net"],r["unrealized"],r["equity"],r["remaining_qty"],stop["realized"],stop["unrealized"],stop["equity"]]
    for col, x in enumerate(expected, 1):
        cell("10月27日デイ途中", f"{openpyxl.utils.get_column_letter(col)}{n}", x)
    formula("10月27日デイ途中", f"E{n}", f"=B{n}-C{n}+D{n}")
    formula("10月27日デイ途中", f"G{n}", f"=E{n}+F{n}")
    formula("10月27日デイ途中", f"K{n}", f"=I{n}+J{n}")
expected_tables["10月27日デイ途中"] = ("D154_current", "A5:K28")

# ----------------------------------------
# 全シートの正式テーブル・数式キャッシュ・エラー値を調べる。
for name in names:
    sh = formulas[name]
    table_name, table_ref = expected_tables[name]
    check(len(sh.tables), 1, f"{name}の正式表数")
    check(list(sh.tables), [table_name], f"{name}の正式表名")
    if table_name in sh.tables:
        t = sh.tables[table_name]
        check(t.ref, table_ref, f"{name}の表範囲")
        check(t.autoFilter.ref if t.autoFilter else None, table_ref, f"{name}のフィルター範囲")
        check([c.name for c in t.tableColumns], [c.value for c in sh[5]][:len(t.tableColumns)], f"{name}の表見出し")
    check(sh.freeze_panes, "A6", f"{name}の見出し固定")
    for row in sh:
        for c in row:
            if c.data_type == "f":
                v = values[name][c.coordinate]
                check(v.data_type == "e", False, f"{name}!{c.coordinate}式エラー")
                check(v.value is not None, True, f"{name}!{c.coordinate}計算値保存")
            check(c.data_type == "e", False, f"{name}!{c.coordinate}セルエラー")
with zipfile.ZipFile(BOOK) as z:
    paths = [n for n in z.namelist() if n.startswith("xl/tables/") and n.endswith(".xml")]
    check(len(paths), 25, "ZIP内の正式表定義数")
    for p in paths:
        t = ET.fromstring(z.read(p))
        check(t.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}autoFilter") is not None, True, f"{p}の自動フィルター定義")
check(sha(BOOK), before, "読み取り検査前後の指紋")
check(sha(COPY), before, "公開版とoutputs版の内容一致")

# ----------------------------------------
# 全25枚を目視できる一覧画像を作る。元Excelには書き込まない。
font = ImageFont.truetype("C:/Windows/Fonts/meiryo.ttc", 20)
sheet_images = [QA / f"{name}.png" for name in names]
contacts = []
for first in range(0, len(sheet_images), 4):
    selected = sheet_images[first:first+4]
    canvas = Image.new("RGB", (2200, 1900), "#e5eaf0")
    draw = ImageDraw.Draw(canvas)
    for j, p in enumerate(selected):
        im = Image.open(p).convert("RGB")
        im.thumbnail((1080, 900))
        x, y = (j%2)*1100+10, (j//2)*950+40
        canvas.paste(im, (x,y))
        draw.text((x,y-32), p.stem, font=font, fill="#152e45")
    target = QA / f"全シート確認_{first//4+1:02}.png"
    canvas.save(target)
    contacts.append(str(target))
result = {"passed":not issues, "checks":dict(counts), "issues":issues, "book_sha256":before,
          "report_sha256":sha(WORK/"report154_combined.json"), "sheets":names, "tables":expected_tables,
          "contacts":contacts, "visual_review":"別途全25描画を目視して追記する", "native_excel_recalculation":False,
          "excel_modified":False}
(QA/"保存Excelの独立検算154.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps({"passed":not issues,"checks":dict(counts),"issues":issues[:10]},ensure_ascii=False))
if issues:
    raise SystemExit(1)

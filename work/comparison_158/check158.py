"""保存したExcelを読み取り専用で検算する。"""

from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import hashlib
import json
import zipfile
import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
BOOK = ROOT / "アルゴリズム調査/00158_D07と改修候補の比較と未解決点/D07と改修候補の日中損益.xlsx"


def digest(path: Path) -> str:
    """機能: 保存ファイルの指紋を求める。
    引数: path (Path) はファイル。
    返り値: str はSHA-256。
    """
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    """機能: 値・数式・正式表・対象日を別実装で照合する。
    引数: なし。
    返り値: None。Excel本体は変更しない。
    """
    report_path = HERE / "report158.json"
    report = json.loads(report_path.read_text(encoding="utf-8"))
    qa = json.loads((HERE / "xlsx_checks/Excel照合158.json").read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(BOOK, data_only=False)
    values = openpyxl.load_workbook(BOOK, data_only=True)
    checks, formulas = 0, 0

    def equal(actual, expected, label):
        """機能: 数値と文字を比較する。
        引数: actual/expected (object) は両値、label (str) は対象。
        返り値: None。不一致なら停止する。
        """
        nonlocal checks
        checks += 1
        if isinstance(expected, (float, int)):
            assert isinstance(actual, (float, int)) and abs(actual - expected) < 1e-6, (label, actual, expected)
        else:
            assert actual == expected, (label, actual, expected)

    expected_sheets = ["比較"] + ["日経_" + m["id"] for m in report["methods"]]
    equal(wb.sheetnames, expected_sheets, "シート名")
    for name, address, expected in qa["checks"]:
        equal(values[name][address].value, expected, f"{name}/{address}")
    for mi, method in enumerate(report["methods"]):
        sn = "日経_" + method["id"]
        s = wb[sn]
        equal(s.freeze_panes, "A6", sn + "/固定")
        equal(len(s.tables), 1, sn + "/正式表数")
        t = s.tables[list(s.tables)[0]]
        equal(t.ref, "A5:H31", sn + "/表範囲")
        assert t.autoFilter is not None
        for i, row in enumerate(method["rows"]):
            n = i + 6
            equal(values[sn][f"A{n}"].value, row["session"], sn + "/場")
            equal(values[sn][f"B{n}"].value, row["direction"], sn + "/方向")
            equal(s[f"E{n}"].value, f"=C{n}-D{n}", sn + "/収支式")
            equal(s[f"F{n}"].value, f"=E{n}" if i == 0 else f"=F{n-1}+E{n}", sn + "/累計式")
        n = mi + 6
        equal(values["比較"][f"A{n}"].value, method["id"], "比較ID")
        equal(wb["比較"][f"D{n}"].value, f"='{sn}'!E32", "比較収支参照")
    for sheet in wb:
        equal(sheet.sheet_view.showGridLines, False, sheet.title + "/目盛")
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    formulas += 1
                    v = values[sheet.title][cell.coordinate]
                    assert v.data_type != "e" and isinstance(v.value, (float, int)), (sheet.title, cell.coordinate, v.value)
                assert cell.data_type != "e", (sheet.title, cell.coordinate)
    with zipfile.ZipFile(BOOK) as z:
        tables = [n for n in z.namelist() if n.startswith("xl/tables/table") and n.endswith(".xml")]
        equal(len(tables), 7, "ZIP正式表数")
    equal(digest(BOOK), qa["book_sha256"], "Excel指紋")
    equal(digest(report_path), qa["report_sha256"], "元会計指紋")
    equal(digest(BOOK), digest(ROOT / "outputs/00158_comparison" / BOOK.name), "出力複製")
    result = {"passed": True, "checked_at": datetime.now(ZoneInfo("Asia/Tokyo")).isoformat(),
              "checks": checks, "formula_count": formulas, "sheets": 7, "tables": 7,
              "complete_session_rows": 156, "financial_replays": 0,
              "book_sha256": digest(BOOK), "report_sha256": digest(report_path),
              "native_excel_recalculation": False, "book_modified": False}
    (HERE / "xlsx_checks/保存Excelの独立検算158.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()

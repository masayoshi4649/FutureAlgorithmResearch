"""保存済みExcelと独立集計を読み取り専用で照合し、監査資料だけを作成する。"""
from pathlib import Path
import hashlib
import json
import math
import statistics
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
import openpyxl
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
REPORT = HERE.parent / "report153_combined.json"
BOOK = ROOT / "アルゴリズム調査/00153_移動平均の傾きと比例枚数/移動平均の傾きと比例枚数の全セッション損益.xlsx"
issues = []
counts = Counter()


def check(actual, expected, where):
    """値を独立期待と照合する。

    引数: actual/expected(object)、where(str)。返り値: なし。不一致をissuesへ記録。
    """
    counts["比較した値"] += 1
    if isinstance(expected, (float, int)) and not isinstance(expected, bool):
        okay = isinstance(actual, (float, int)) and math.isfinite(actual) and abs(actual - expected) <= 1e-5
    else:
        okay = actual == expected
    if not okay:
        issues.append({"対象": where, "保存値": actual, "期待": expected})


def hashfile(path):
    """対象の内容指紋を取得する。

    引数: path(Path)。返り値: strはSHA256。
    """
    return hashlib.sha256(path.read_bytes()).hexdigest()


def value(sheet, cell, expected):
    """保存済み計算値を比較する。

    引数: sheet(str)、cell(str)、expected(object)。返り値: なし。
    """
    check(values[sheet][cell].value, expected, f"{sheet}!{cell}")


def formula(sheet, cell, expected):
    """保存数式を期待と完全一致で比較する。

    引数: sheet(str)、cell(str)、expected(str)。返り値: なし。
    """
    check(formulas[sheet][cell].value, expected, f"{sheet}!{cell}の式")
    counts["照合した数式"] += 1


before = hashfile(BOOK)
report = json.loads(REPORT.read_text(encoding="utf-8"))
values = openpyxl.load_workbook(BOOK, data_only=True)
formulas = openpyxl.load_workbook(BOOK, data_only=False)
methods = report["methods"]
check(len(methods), 68, "手法組数")
expected_names = ["比較"] + [("日経" if m["symbol"] == "NK225" else "TOPIX") + "_" + m["id"] for m in methods] + ["10月26日ナイト途中"]
check(values.sheetnames, expected_names, "全シート名と順番")
check(len(values.sheetnames), 70, "シート数")

# ----------------------------------------
# 保存済み明細と数式を、JSONの全51場および独立再集計に照合する。
for m in methods:
    name = ("日経" if m["symbol"] == "NK225" else "TOPIX") + "_" + m["id"]
    rows = sorted(m["rows"], key=lambda row: row["session"])
    check(len(rows), 51, name + "場数")
    check(len({row["session"] for row in rows}), 51, name + "一意場数")
    cumulative = 0
    for n, row in enumerate(rows, 6):
        cumulative += row["net"]
        value(name, f"A{n}", row["session"])
        value(name, f"B{n}", "デイ" if row["kind"] == "Day" else "ナイト")
        for col, key in {"C": "profit", "D": "loss", "G": "parent_profit", "H": "parent_loss", "I": "entries", "J": "additions", "K": "reductions", "L": "max_qty", "M": "held_rate", "N": "equity_drawdown"}.items():
            value(name, f"{col}{n}", row[key])
        value(name, f"E{n}", row["profit"] - row["loss"])
        value(name, f"F{n}", cumulative)
        formula(name, f"E{n}", f"=C{n}-D{n}")
        formula(name, f"F{n}", "=E6" if n == 6 else f"=F{n-1}+E{n}")
    for col, key in {"C": "profit", "D": "loss", "G": "parent_profit", "H": "parent_loss", "I": "entries", "J": "additions", "K": "reductions"}.items():
        value(name, f"{col}57", sum(row[key] for row in rows))
        formula(name, f"{col}57", f"=SUM({col}6:{col}56)")
    value(name, "E57", sum(r["profit"] - r["loss"] for r in rows))
    value(name, "F57", cumulative)
    value(name, "L57", max(r["max_qty"] for r in rows))
    value(name, "M57", statistics.mean(r["held_rate"] for r in rows))
    value(name, "N57", max(r["equity_drawdown"] for r in rows))
    for cell, expected in {"E57": "=C57-D57", "F57": "=F56", "L57": "=MAX(L6:L56)", "M57": "=AVERAGE(M6:M56)", "N57": "=MAX(N6:N56)"}.items():
        formula(name, cell, expected)
    counts["完了場行"] += len(rows)

# ----------------------------------------
# 比較は元summaryだけに依存せず、51場の値から再計算する。
for n, m in enumerate(methods, 6):
    name = ("日経" if m["symbol"] == "NK225" else "TOPIX") + "_" + m["id"]
    rows = m["rows"]
    profit = sum(r["profit"] for r in rows)
    loss = sum(r["loss"] for r in rows)
    net = profit - loss
    losses = [-r["net"] for r in rows if r["net"] < 0]
    positive = sorted((r["net"] for r in rows if r["net"] > 0), reverse=True)
    expected = {"A": name, "B": m["description"], "C": profit, "D": loss, "E": net, "F": profit / loss if loss else None, "G": sum(r["net"] > 0 for r in rows), "H": len(losses), "I": sum(r["entries"] == 0 for r in rows), "J": max(losses, default=0), "K": statistics.mean(losses) if losses else None, "L": statistics.median(losses) if losses else 0, "M": net - sum(positive[:3]), "N": sum(r["net"] for r in rows if r["kind"] == "Day"), "O": sum(r["net"] for r in rows if r["kind"] == "Night"), "P": sum(r["entries"] for r in rows)}
    for col, wanted in expected.items():
        value("比較", f"{col}{n}", wanted)
    prefix = f"'{name}'!"
    nets, kinds = prefix + "$E$6:$E$56", prefix + "$B$6:$B$56"
    expected_f = {"C": f"={prefix}C57", "D": f"={prefix}D57", "E": f"={prefix}E57", "F": f'=IF(D{n}=0,"",C{n}/D{n})', "G": f'=COUNTIFS({nets},">0")', "H": f'=COUNTIFS({nets},"<0")', "I": f"=COUNTIFS({prefix}$I$6:$I$56,0)", "J": f"=MAX(0,-MIN({nets}))", "K": f'=IF(H{n}=0,"",-SUMIFS({nets},{nets},"<0")/H{n})', "M": f"=E{n}-MAX(0,LARGE({nets},1))-MAX(0,LARGE({nets},2))-MAX(0,LARGE({nets},3))", "N": f'=SUMIFS({nets},{kinds},"デイ")', "O": f'=SUMIFS({nets},{kinds},"ナイト")', "P": f"={prefix}I57"}
    for col, wanted in expected_f.items():
        formula("比較", f"{col}{n}", wanted)
    counts["比較行"] += 1

check(len(report["partial"]), 68, "途中行数")
for n, row in enumerate(report["partial"], 6):
    name = ("日経" if row["symbol"] == "NK225" else "TOPIX") + "_" + row["id"]
    for col, wanted in {"A": name, "B": row["profit"], "C": row["loss"], "D": row["profit"]-row["loss"], "E": row["unrealized"], "F": row["net"]+row["unrealized"], "G": row["entries"], "H": row["remaining_qty"]}.items():
        value("10月26日ナイト途中", f"{col}{n}", wanted)
    formula("10月26日ナイト途中", f"D{n}", f"=B{n}-C{n}")
    formula("10月26日ナイト途中", f"F{n}", f"=D{n}+E{n}")
    counts["途中行"] += 1

# ----------------------------------------
# 全シートの正式テーブルとフィルター・見出し・数式エラーを検査する。
table_records = []
for sheet in formulas:
    check(len(sheet.tables), 1, sheet.title + "正式テーブル数")
    check(sheet.freeze_panes, "A6", sheet.title + "見出し固定")
    expected_ref = "A5:P73" if sheet.title == "比較" else "A5:H73" if sheet.title == "10月26日ナイト途中" else "A5:N56"
    for table in sheet.tables.values():
        check(table.ref, expected_ref, sheet.title + "テーブル範囲")
        check(table.autoFilter is not None, True, sheet.title + "フィルターあり")
        check(table.autoFilter.ref, expected_ref, sheet.title + "フィルター範囲")
        for n, col in enumerate(table.tableColumns, 1):
            check(col.name, sheet.cell(5, n).value, sheet.title + f"列見出し{n}")
        table_records.append({"シート": sheet.title, "テーブル": table.name, "範囲": table.ref, "フィルター": table.autoFilter.ref})
    for row in sheet:
        for cell in row:
            cached = values[sheet.title][cell.coordinate]
            if cached.data_type == "e":
                issues.append({"対象": f"{sheet.title}!{cell.coordinate}", "エラー": cached.value})
            if cell.data_type == "f":
                counts["全保存数式"] += 1
                if cached.value is None:
                    issues.append({"対象": f"{sheet.title}!{cell.coordinate}", "エラー": "数式キャッシュなし"})
                if any(s in cell.value for s in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")):
                    issues.append({"対象": f"{sheet.title}!{cell.coordinate}", "エラー": "数式内エラー"})
    counts["正式テーブル"] += len(sheet.tables)
with zipfile.ZipFile(BOOK) as z:
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    check(len([p for p in z.namelist() if p.startswith("xl/tables/table") and p.endswith(".xml")]), 70, "ZIP内正式テーブル数")
    for path in [p for p in z.namelist() if p.startswith("xl/worksheets/sheet") and p.endswith(".xml")]:
        tree = ET.fromstring(z.read(path))
        check(len(tree.findall("s:tableParts/s:tablePart", ns)), 1, path + "テーブル関連付け")

# ----------------------------------------
# 元PNGを改変せず、全68手法画像のコンタクトシートだけを監査用に作る。
font = ImageFont.truetype("C:/Windows/Fonts/meiryo.ttc", 22)
method_images = [HERE / (name + ".png") for name in expected_names[1:-1]]
contacts = []
for offset in range(0, len(method_images), 6):
    batch = method_images[offset:offset+6]
    tile_w, tile_h = 936, 970
    canvas = Image.new("RGB", (tile_w*2, tile_h*3), "#eef2f7")
    draw = ImageDraw.Draw(canvas)
    for n, path in enumerate(batch):
        picture = Image.open(path).convert("RGB")
        picture.thumbnail((tile_w, tile_h-32))
        x, y = (n%2)*tile_w, (n//2)*tile_h
        draw.text((x+10, y+2), path.stem, font=font, fill="#172c42")
        canvas.paste(picture, (x, y+32))
    target = HERE / f"監査一覧_{offset//6+1:02d}.png"
    canvas.save(target)
    contacts.append({"画像": target.name, "元画像": [p.name for p in batch]})

check(hashfile(BOOK), before, "監査前後のxlsx内容指紋")
out = {"対象": str(BOOK), "xlsx_sha256": before, "report_sha256": hashfile(REPORT), "保存ファイル変更": False, "読取手段": "openpyxl値/数式の別読込とZIP内XML。xlsxの保存・再計算・再出力なし。", "計数": dict(counts), "不一致": issues, "テーブル": table_records, "視覚確認用": contacts, "視覚確認": "コンタクトシートおよび比較/途中画像を別途全件目視する。", "限界": "Excelネイティブアプリでの再計算や並べ替え操作は未実施。保存数式の参照と保存計算値を元会計へ全件照合。"}
(HERE / "独立検算153.json").write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"計数": dict(counts), "不一致件数": len(issues), "不一致先頭": issues[:5], "xlsx_sha256": before}, ensure_ascii=False))

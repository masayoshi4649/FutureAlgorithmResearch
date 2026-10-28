"""保存Excelを読み取り専用で検算し、正式Tableと全シート描画の確認資料を作る。"""

from collections import Counter
from pathlib import Path
import hashlib
import json
import math
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
from PIL import Image, ImageDraw, ImageFont


WORK = Path(__file__).resolve().parent
ROOT = WORK.parents[1]
QA = WORK/"xlsx_checks"
NAME = "利益保護と再参加と入口見直しの全日中損益.xlsx"
BOOK = ROOT/"アルゴリズム調査/00156_固定値幅の利益保護と遅い入口の再検証"/NAME
COPY = ROOT/"outputs/00156_s01_exit"/NAME
issues = []
counts = Counter()


def sha(path: Path) -> str:
    """内容指紋を読み取りだけで求める。

    引数: path(Path)はファイル。返り値: strはSHA256。
    """
    return hashlib.sha256(path.read_bytes()).hexdigest()


def check(actual, expected, label: str) -> None:
    """保存値と独立会計の期待値を比較する。

    引数: actual/expected(object)は値、label(str)は対象。返り値: None。
    """
    counts["検査数"] += 1
    okay = actual == expected
    if isinstance(expected,(int,float)) and not isinstance(expected,bool):
        okay = isinstance(actual,(int,float)) and math.isfinite(actual) and abs(actual-expected)<=1e-5
    if not okay:
        issues.append({"対象":label,"保存値":actual,"期待":expected})


def cell(sheet: str, address: str, expected) -> None:
    """保存された計算済み値又は文字を照合する。

    引数: sheet/address(str)は位置、expected(object)は期待値。返り値: None。
    """
    check(values[sheet][address].value,expected,f"{sheet}!{address}")


def formula(sheet: str, address: str, expected: str) -> None:
    """数式文字列を完全一致で検査する。

    引数: sheet/address/expected(str)は位置と式。返り値: None。
    """
    check(formulas[sheet][address].value,expected,f"{sheet}!{address}の式")
    counts["数式照合"] += 1


before = sha(BOOK)
report = json.loads((WORK/"report156_all.json").read_text(encoding="utf-8"))
values = openpyxl.load_workbook(BOOK,data_only=True)
formulas = openpyxl.load_workbook(BOOK,data_only=False)
methods = report["methods"]
names = ["比較"]+[f'日経_{m["id"]}' for m in methods]+["10月27日デイ途中"]
check(values.sheetnames,names,"シート順序")
check(len(methods),44,"最終候補44手法")
last = len(methods)+5
tables = {"比較":("D156_comparison",f"A5:M{last}"),"10月27日デイ途中":("D156_current",f"A5:K{last}")}

# ----------------------------------------
# 手法ごとの26完了場と、収支・累計・時価最悪を照合する。
for method in methods:
    name=f'日経_{method["id"]}'
    rows=sorted(method["rows"],key=lambda r:r["session"])
    check(len(rows),26,name+"完了場数")
    check(len({r["session"] for r in rows}),26,name+"重複なし")
    cumulative=0
    for n,row in enumerate(rows,6):
        cumulative+=row["net"]
        risk=row["intraday_risk"]
        data=[row["session"],row["direction"],row["profit"],row["loss"],row["net"],cumulative,
              risk["min_equity"],risk["max_drawdown"]]
        for c,value in enumerate(data,1):
            cell(name,f"{openpyxl.utils.get_column_letter(c)}{n}",value)
        formula(name,f"E{n}",f"=C{n}-D{n}")
        formula(name,f"F{n}",f"=E{n}" if n==6 else f"=F{n-1}+E{n}")
    for col,key in (("C","profit"),("D","loss"),("E","net"),("F","net")):
        cell(name,col+"32",method["summary"][key])
    cell(name,"G32",min(r["intraday_risk"]["min_equity"] for r in rows))
    cell(name,"H32",max(r["intraday_risk"]["max_drawdown"] for r in rows))
    for col in ("C","D"):
        formula(name,col+"32",f"=SUM({col}6:{col}31)")
    formula(name,"E32","=C32-D32")
    formula(name,"F32","=F31")
    formula(name,"G32","=MIN(G6:G31)")
    formula(name,"H32","=MAX(H6:H31)")
    tables[name]=(f'D156_{method["id"]}',"A5:H31")
    counts["完了場の掲載行"]+=26

# ----------------------------------------
# 比較表と途中を保存会計に全件一致させる。
for n,method in enumerate(methods,6):
    s=method["summary"]
    p=f'\'日経_{method["id"]}\'!'
    net=f"{p}$E$6:$E$31"
    category="基準" if method["id"]=="S01" else method.get("group_label", "再参加の変更" if method["id"].startswith("R") else "再参加の損切り変更" if method["id"].startswith("L") else "入口・出口の変更")
    display=values["比較"][f"B{n}"].value
    check(isinstance(display,str) and display.startswith(category+"："),True,"比較の変更区分と説明")
    check("core入口E" in display,False,"入口条件を調査IDだけで説明しない")
    data=[method["id"],display,s["profit"],s["loss"],s["net"],s["pf"] if s["loss"] else "",
          s["losses"],s["avgloss"],s["maxloss"],s["ex3"],s["entries"],
          min(r["intraday_risk"]["min_equity"] for r in method["rows"]),max(r["intraday_risk"]["max_drawdown"] for r in method["rows"])]
    for c,value in enumerate(data,1):
        cell("比較",f"{openpyxl.utils.get_column_letter(c)}{n}",value)
    expr=[f"={p}C32",f"={p}D32",f"={p}E32",f'=IF(D{n}=0,"",C{n}/D{n})',f'=COUNTIFS({net},"<0")',
          f'=IF(G{n}=0,0,-SUMIFS({net},{net},"<0")/G{n})',f"=MAX(0,-MIN({net}))",
          f"=E{n}-MAX(0,LARGE({net},1))-MAX(0,LARGE({net},2))-MAX(0,LARGE({net},3))"]
    for c,value in enumerate(expr,3):
        formula("比較",f"{openpyxl.utils.get_column_letter(c)}{n}",value)
    formula("比較",f"L{n}",f"={p}G32")
    formula("比較",f"M{n}",f"={p}H32")
for n,row in enumerate(report["current_rows"],6):
    r=row["intraday_risk"]
    data=[row["id"],row["profit"],row["loss"],row["open_parent_realized"],row["net"],row["unrealized"],row["equity"],
          row["remaining_qty"],row["entries"],r["min_equity"],r["max_drawdown"]]
    for c,value in enumerate(data,1):
        cell("10月27日デイ途中",f"{openpyxl.utils.get_column_letter(c)}{n}",value)
    formula("10月27日デイ途中",f"E{n}",f"=B{n}-C{n}+D{n}")
    formula("10月27日デイ途中",f"G{n}",f"=E{n}+F{n}")

# ----------------------------------------
# 正式Tableと自動フィルター、数式キャッシュ、計算エラーを全シート検査する。
for name in names:
    sheet=formulas[name]
    table_name,ref=tables[name]
    check(list(sheet.tables),[table_name],name+"正式表名")
    if table_name in sheet.tables:
        t=sheet.tables[table_name]
        check(t.ref,ref,name+"表範囲")
        check(t.autoFilter.ref if t.autoFilter else None,ref,name+"フィルター")
        check([c.name for c in t.tableColumns],[c.value for c in sheet[5]][:len(t.tableColumns)],name+"表見出し")
    check(sheet.freeze_panes,"A6",name+"見出し固定")
    for row in sheet:
        for c in row:
            check(c.data_type=="e",False,name+"!"+c.coordinate+"エラーなし")
            if c.data_type=="f":
                result=values[name][c.coordinate]
                check(result.data_type=="e",False,name+"!"+c.coordinate+"式エラーなし")
                check(result.value is not None,True,name+"!"+c.coordinate+"計算キャッシュ")
with zipfile.ZipFile(BOOK) as archive:
    definitions=[p for p in archive.namelist() if p.startswith("xl/tables/") and p.endswith(".xml")]
    check(len(definitions),len(names),"XML正式表定義数")
    for filename in definitions:
        node=ET.fromstring(archive.read(filename))
        check(node.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}autoFilter") is not None,True,filename+"フィルターXML")
check(sha(BOOK),before,"読取検査で成果未変更")
check(sha(COPY),before,"公開とoutputs同一")

# ----------------------------------------
# 全描画を見比べる一覧画像を作り、Excel自体は一切書換えない。
font=ImageFont.truetype("C:/Windows/Fonts/meiryo.ttc",20)
contacts=[]
for first in range(0,len(names),4):
    canvas=Image.new("RGB",(2200,1900),"#e5eaf0")
    draw=ImageDraw.Draw(canvas)
    for j,name in enumerate(names[first:first+4]):
        image=Image.open(QA/(name+".png")).convert("RGB")
        image.thumbnail((1080,900))
        x,y=(j%2)*1100+10,(j//2)*950+40
        canvas.paste(image,(x,y))
        draw.text((x,y-32),name,font=font,fill="#152e45")
    target=QA/f"全シート確認_{first//4+1:02}.png"
    canvas.save(target)
    contacts.append(str(target))
out={"passed":not issues,"checks":dict(counts),"issues":issues,"book_sha256":before,
     "report_sha256":sha(WORK/"report156_all.json"),"sheet_names":names,"tables":tables,"contacts":contacts,
     "visual_review":"全46描画を別途確認","native_excel_recalculation":False,"excel_modified":False}
(QA/"保存Excelの独立検算156.json").write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps({"passed":not issues,"checks":dict(counts),"issues":issues[:10]},ensure_ascii=True))
if issues:
    raise SystemExit(1)
